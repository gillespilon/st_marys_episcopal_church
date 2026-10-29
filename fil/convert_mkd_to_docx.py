#!/usr/bin/env python3
import os
import re
from PIL import Image as PILImage
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_hyperlink(paragraph, url, text, font_name="Liberation Sans", color="0000FF", underline=True):
    """Appends an active hyperlink with explicit text runs and styling."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    # Explicit font family on hyperlink run
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rPr.append(rFonts)

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)

    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def parse_markdown_paragraph(paragraph, text):
    """Parses paragraph text for [label](url) links."""
    link_pattern = re.compile(r'\[(.*?)\]\((.*?)\)')
    last_idx = 0

    for match in link_pattern.finditer(text):
        start, end = match.span()
        label, url = match.group(1), match.group(2)

        if start > last_idx:
            paragraph.add_run(text[last_idx:start])

        if url.endswith('.html') and not url.startswith(('http://', 'https://')):
            url = re.sub(r'\.html$', '.docx', url)

        add_hyperlink(paragraph, url, label)
        last_idx = end

    if last_idx < len(text):
        paragraph.add_run(text[last_idx:])

def apply_custom_document_styles(doc):
    """Applies global document formatting."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Liberation Sans'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.line_spacing = 1.15

    h1_style = doc.styles['Heading 1']
    h1_font = h1_style.font
    h1_font.name = 'Liberation Sans'
    h1_font.size = Pt(18)
    h1_font.bold = True
    h1_font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(6)

    h2_style = doc.styles['Heading 2']
    h2_font = h2_style.font
    h2_font.name = 'Liberation Sans'
    h2_font.size = Pt(14)
    h2_font.bold = True
    h2_font.color.rgb = RGBColor(0x00, 0x55, 0x88)
    h2_style.paragraph_format.space_before = Pt(10)
    h2_style.paragraph_format.space_after = Pt(4)

    h3_style = doc.styles['Heading 3']
    h3_font = h3_style.font
    h3_font.name = 'Liberation Sans'
    h3_font.size = Pt(12)
    h3_font.bold = True
    h3_font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    h3_style.paragraph_format.space_before = Pt(8)
    h3_style.paragraph_format.space_after = Pt(2)

def convert_mkd_file(mkd_path, docx_path, image_dirs=["fot", "../fot"]):
    doc = Document()
    apply_custom_document_styles(doc)

    max_width_in = 6.5  # Max printable width within 1-inch margins on 8.5" page

    with open(mkd_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
            i += 1
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
            i += 1
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
            i += 1

        elif line.startswith("<table"):
            table_lines = []
            while i < len(lines):
                table_lines.append(lines[i])
                if "</table>" in lines[i]:
                    i += 1
                    break
                i += 1

            table_html = "\n".join(table_lines)
            soup = BeautifulSoup(table_html, 'html.parser')
            rows = soup.find_all('tr')

            if rows:
                num_rows = len(rows)
                num_cols = max(len(r.find_all(['td', 'th'])) for r in rows)
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.style = 'Table Grid'

                for r_idx, row in enumerate(rows):
                    cells = row.find_all(['td', 'th'])
                    for c_idx, cell in enumerate(cells):
                        table.cell(r_idx, c_idx).text = cell.get_text().strip()

        elif line.startswith("!") and "(" in line and ")" in line:
            img_src = line.split("(")[1].split(")")[0]
            img_filename = os.path.basename(img_src)

            found_path = None
            for d in image_dirs:
                candidate = os.path.join(d, img_filename)
                if os.path.exists(candidate):
                    found_path = candidate
                    break

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # FIX WHITE HORIZONTAL LINES: Set exact/single line-spacing on image paragraphs
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

            if found_path:
                # FIX IMAGE RESIZING: Maintain native DPI/dimensions up to max page width
                with PILImage.open(found_path) as img:
                    width_px, _ = img.size
                    dpi = img.info.get('dpi', (96, 96))[0]
                    width_in = width_px / float(dpi if dpi > 0 else 96)

                run = p.add_run()
                if width_in > max_width_in:
                    run.add_picture(found_path, width=Inches(max_width_in))
                else:
                    run.add_picture(found_path, width=Inches(width_in))
            else:
                p.add_run(f"[Image Missing: {img_filename}]")
            i += 1

        else:
            p = doc.add_paragraph()
            parse_markdown_paragraph(p, line)
            i += 1

    doc.save(docx_path)
    print(f"Created: {docx_path}")

if __name__ == "__main__":
    fil_dir = os.path.dirname(os.path.abspath(__file__))
    repo_root = os.path.abspath(os.path.join(fil_dir, ".."))

    image_search_dirs = [
        os.path.join(repo_root, "fot"),
        os.path.join(repo_root, "fol"),
        os.path.join(repo_root, "fov", "thrift_shop"),
    ]

    print(f"Scanning for .mkd files in: {repo_root}")

    count = 0
    for root, _, files in os.walk(repo_root):
        for file in files:
            if file.endswith(".mkd"):
                src_path = os.path.join(root, file)
                dst_path = os.path.splitext(src_path)[0] + ".docx"

                print(f"Converting: {src_path} -> {dst_path}")
                convert_mkd_file(src_path, dst_path, image_dirs=image_search_dirs)
                count += 1

    print(f"\nDone! Successfully converted {count} file(s).")
